import io
import re
from typing import List, Dict, Any
import pypdf
import docx
from .sensitive_detector import PATTERNS, CATEGORY_ROLE_MAP

# High-Precision Secret and Credential Patterns
SECRET_PATTERNS = {
    "OpenAI / Anthropic API Key": re.compile(
        r"\b(?:sk-[a-zA-Z0-9_\-]{20,}|sk-proj-[a-zA-Z0-9_\-]{20,}|sk-ant-[a-zA-Z0-9_\-]{20,})\b"
    ),
    "AWS Access Key ID": re.compile(
        r"\b(AKIA[0-9A-Z]{16})\b"
    ),
    "AWS Secret Access Key": re.compile(
        r"(?i)aws_secret_access_key\s*[:=]\s*['\"]?([A-Za-z0-9\/+=]{40})['\"]?"
    ),
    "GitHub Personal Access Token": re.compile(
        r"\b(?:ghp_[a-zA-Z0-9]{36}|gho_[a-zA-Z0-9]{36}|github_pat_[a-zA-Z0-9_]{82})\b"
    ),
    "Slack Token / Webhook": re.compile(
        r"(?:xox[baprs]-[0-9a-zA-Z]{10,48}|https:\/\/hooks\.slack\.com\/services\/T[0-9A-Z]{8}\/B[0-9A-Z]{8}\/[0-9a-zA-Z]{24})"
    ),
    "Google Cloud / Firebase API Key": re.compile(
        r"\bAIza[0-9A-Za-z\-_]{35}\b"
    ),
    "Database Connection String": re.compile(
        r"\b(?:postgres|postgresql|mysql|mongodb|mongodb\+srv|redis):\/\/[^:\s]+:[^@\s]+@[a-zA-Z0-9_\-\.]+:[0-9]{2,5}(?:\/[a-zA-Z0-9_\-]+)?\b",
        re.IGNORECASE
    ),
    "Private Key Block": re.compile(
        r"-----BEGIN (?:RSA |EC |OPENSSH |DSA )?PRIVATE KEY-----"
    ),
    "JWT Bearer Token": re.compile(
        r"\beyJ[a-zA-Z0-9_\-]+\.eyJ[a-zA-Z0-9_\-]+\.[a-zA-Z0-9_\-]+\b"
    ),
    "Generic High-Entropy Secret Assignment": re.compile(
        r"(?i)(?:api_key|apikey|secret_key|secret|password|pwd|auth_token|access_token|client_secret)\s*[:=]\s*['\"]([a-zA-Z0-9_\-\.\$\!\@\#\%\^\&\*]{8,})['\"]"
    ),
}

# Contextual Keywords in Code and Documents
CONTEXT_HINTS = {
    "Credentials & Passwords": re.compile(r"(?i)\b(password|passwd|pwd|secret|credential|token)\s*[:=]"),
    "Financial & Payment": re.compile(r"(?i)\b(salary|bank|credit\s*card|debit\s*card|cvv|routing|iban|swift|account\s*number)\b"),
    "Government Identification": re.compile(r"(?i)\b(ssn|social\s*security|national\s*id|passport|aadhaar|pan\s*card|driving\s*licen[sc]e)\b"),
    "Healthcare & Medical": re.compile(r"(?i)\b(diagnosis|patient|prescription|treatment|blood\s*group|disease|medical\s*record)\b"),
}


def mask_evidence(text: str, keep_chars: int = 3) -> str:
    """Masks secret or sensitive snippet in report so raw secrets aren't leaked in findings."""
    clean = str(text).strip()
    if len(clean) <= keep_chars * 2:
        return "*" * len(clean)
    return clean[:keep_chars] + "*" * (len(clean) - keep_chars * 2) + clean[-keep_chars:]


def extract_text_from_pdf(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extracts text page by page from PDF bytes."""
    pages_content = []
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages_content.append({
                "page": page_idx + 1,
                "text": text
            })
    except Exception as e:
        pages_content.append({"page": 1, "text": f"[Error reading PDF: {str(e)}]"})
    return pages_content


def extract_text_from_docx(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extracts text paragraph by paragraph and tables from DOCX."""
    doc = docx.Document(io.BytesIO(file_bytes))
    sections = []
    
    # Paragraphs
    for p_idx, p in enumerate(doc.paragraphs):
        if p.text.strip():
            sections.append({
                "location": f"Paragraph {p_idx + 1}",
                "text": p.text
            })
            
    # Tables
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                sections.append({
                    "location": f"Table {t_idx + 1}, Row {r_idx + 1}",
                    "text": " | ".join(row_texts)
                })
    return sections


def scan_text_content(content_blocks: List[Dict[str, Any]], source_name: str) -> List[Dict[str, Any]]:
    """
    Deep scan across document text or code blocks:
    - Secret patterns (API keys, Tokens, DB Strings, Private Keys)
    - PII patterns (Emails, Phones, SSNs, Cards, IPv4, Government IDs)
    - Contextual keyword triggers (Passwords, medical records, financial markers)
    """
    findings = []
    seen_matches = set()

    for block in content_blocks:
        location = block.get("location") or f"Page {block.get('page', 1)}"
        text = block.get("text", "")

        lines = text.split("\n")
        for line_num, line in enumerate(lines, 1):
            stripped_line = line.strip()
            if not stripped_line:
                continue

            # 1. High-Priority Secret Patterns
            for secret_type, secret_regex in SECRET_PATTERNS.items():
                for match in secret_regex.finditer(stripped_line):
                    matched_val = match.group(0)
                    key = (secret_type, matched_val, location, line_num)
                    if key not in seen_matches:
                        seen_matches.add(key)
                        findings.append({
                            "Source": source_name,
                            "Location": f"{location}, Line {line_num}",
                            "Category": "Authentication Secret",
                            "Role": "Direct Secret",
                            "Detection Method": f"Secret Token Rule ({secret_type})",
                            "Evidence": mask_evidence(matched_val, keep_chars=3),
                            "Confidence": "Very High",
                            "Recommended Action": "Rotate this credential immediately and purge from git history / configuration."
                        })

            # 2. PII Patterns
            for pat_name, pat_regex in PATTERNS.items():
                for match in pat_regex.finditer(stripped_line):
                    matched_val = match.group(0)
                    key = (pat_name, matched_val, location, line_num)
                    if key not in seen_matches:
                        seen_matches.add(key)
                        findings.append({
                            "Source": source_name,
                            "Location": f"{location}, Line {line_num}",
                            "Category": pat_name,
                            "Role": CATEGORY_ROLE_MAP.get(pat_name, "Direct Identifier"),
                            "Detection Method": "Pattern Syntax Match",
                            "Evidence": mask_evidence(matched_val, keep_chars=2),
                            "Confidence": "High",
                            "Recommended Action": "Redact or mask value before distributing this document/code."
                        })

            # 3. Contextual Keyword Findings (e.g. "password = ...", "SSN: ...")
            for ctx_category, ctx_regex in CONTEXT_HINTS.items():
                if ctx_regex.search(stripped_line):
                    # Check if already caught by PII or secret
                    if not any(f["Location"] == f"{location}, Line {line_num}" for f in findings):
                        key = (ctx_category, stripped_line[:40], location, line_num)
                        if key not in seen_matches:
                            seen_matches.add(key)
                            findings.append({
                                "Source": source_name,
                                "Location": f"{location}, Line {line_num}",
                                "Category": ctx_category,
                                "Role": "Contextual Sensitive",
                                "Detection Method": "Contextual Proximity Heuristic",
                                "Evidence": mask_evidence(stripped_line[:50], keep_chars=5),
                                "Confidence": "Medium-High",
                                "Recommended Action": "Review line contents for confidential disclosures."
                            })

    return findings


def scan_document(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Entrypoint to scan any supported document/code/config file."""
    lower_name = filename.lower()
    
    if lower_name.endswith(".pdf"):
        blocks = extract_text_from_pdf(file_bytes)
        return scan_text_content(blocks, filename)
    elif lower_name.endswith(".docx"):
        blocks = extract_text_from_docx(file_bytes)
        return scan_text_content(blocks, filename)
    else:
        # Plain text, markdown, json, python, yaml, env, logs
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="ignore")

        blocks = [{"location": f"Line {i+1}", "text": line} for i, line in enumerate(text.splitlines())]
        return scan_text_content(blocks, filename)
